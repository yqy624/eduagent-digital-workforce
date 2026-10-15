"""Export the EduAgent Markdown plan into a formatted Word document."""
from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


PRIMARY = RGBColor(0x1F, 0x4E, 0x79)
ACCENT = RGBColor(0x2E, 0x74, 0xB5)
TEXT = RGBColor(0x1F, 0x29, 0x37)


def set_east_asia_font(run_or_style, font_name: str) -> None:
    element = run_or_style.element
    rpr = element.get_or_add_rPr() if hasattr(element, "get_or_add_rPr") else element.rPr
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        element.append(rpr)
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font_name)


def shade_cell(cell, fill: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = TEXT
    set_east_asia_font(normal, "宋体")
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(4)

    for name, size, color, before, after in [
        ("Heading 1", 17, PRIMARY, 18, 8),
        ("Heading 2", 14, PRIMARY, 14, 6),
        ("Heading 3", 12, ACCENT, 10, 4),
        ("Heading 4", 11, TEXT, 8, 3),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        set_east_asia_font(style, "黑体")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_cover(doc: Document) -> None:
    for _ in range(5):
        doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("EduAgent")
    run.bold = True
    run.font.size = Pt(36)
    run.font.color.rgb = PRIMARY
    set_east_asia_font(run, "黑体")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("基于多 Agent 的智慧教育学习运营平台")
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = ACCENT
    set_east_asia_font(run, "黑体")

    doc.add_paragraph()
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = note.add_run("项目计划书")
    run.font.size = Pt(14)
    set_east_asia_font(run, "宋体")
    doc.add_page_break()


def add_paragraph_with_inline_code(doc: Document, text: str, style: str | None = None):
    paragraph = doc.add_paragraph(style=style)
    parts = re.split(r"(`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        run = paragraph.add_run(part[1:-1] if part.startswith("`") and part.endswith("`") else part)
        run.font.name = "Consolas" if part.startswith("`") and part.endswith("`") else "Calibri"
        set_east_asia_font(run, "宋体")
    return paragraph


def split_table_row(line: str) -> list[str]:
    stripped = line.strip().strip("|")
    return [cell.strip() for cell in stripped.split("|")]


def is_separator_row(line: str) -> bool:
    return bool(re.match(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?\s*$", line))


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    max_cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=max_cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, row in enumerate(rows):
        for j in range(max_cols):
            cell = table.cell(i, j)
            value = row[j] if j < len(row) else ""
            cell.text = ""
            paragraph = cell.paragraphs[0]
            run = paragraph.add_run(value)
            run.font.size = Pt(9.5)
            run.font.name = "Calibri"
            set_east_asia_font(run, "宋体")
            if i == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                set_east_asia_font(run, "黑体")
                shade_cell(cell, "1F4E79")
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()


def add_code_block(doc: Document, code_lines: list[str]) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.4)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run("\n".join(code_lines))
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    set_east_asia_font(run, "宋体")


def export_markdown_to_docx(markdown_path: Path, output_path: Path) -> None:
    doc = Document()
    configure_document(doc)
    add_cover(doc)

    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    in_code = False
    code_lines: list[str] = []
    table_rows: list[list[str]] = []

    def flush_table() -> None:
        nonlocal table_rows
        if table_rows:
            add_table(doc, table_rows)
            table_rows = []

    for raw_line in lines:
        line = raw_line.rstrip()
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                flush_table()
                in_code = True
            continue

        if in_code:
            code_lines.append(line)
            continue

        if not stripped:
            flush_table()
            continue

        if stripped.startswith("|") and "|" in stripped[1:]:
            if not is_separator_row(stripped):
                table_rows.append(split_table_row(stripped))
            continue

        flush_table()

        heading = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if heading:
            level = min(len(heading.group(1)), 4)
            text = heading.group(2).strip()
            doc.add_heading(text, level=level)
            continue

        bullet = re.match(r"^[-*]\s+(.+)$", stripped)
        if bullet:
            add_paragraph_with_inline_code(doc, bullet.group(1), style="List Bullet")
            continue

        add_paragraph_with_inline_code(doc, stripped)

    flush_table()
    if code_lines:
        add_code_block(doc, code_lines)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("markdown_path", type=Path)
    parser.add_argument("output_path", type=Path)
    args = parser.parse_args()
    export_markdown_to_docx(args.markdown_path, args.output_path)


if __name__ == "__main__":
    main()
