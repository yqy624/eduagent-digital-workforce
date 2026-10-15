"""Build a richer EduAgent project plan Word document with black text only."""
from __future__ import annotations

import argparse
import math
import random
import re
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_MARKDOWN = ROOT / "EduAgent_项目计划书.md"
DEFAULT_OUTPUT = ROOT / "EduAgent_项目计划书.docx"
ASSET_DIR = ROOT / "assets" / "project-plan-rich"

BLACK = RGBColor(0, 0, 0)
LIGHT_GRAY = "E5E7EB"
MID_GRAY = "D1D5DB"
FONT_REGULAR = Path(r"C:\Windows\Fonts\msyh.ttc")
FONT_BOLD = Path(r"C:\Windows\Fonts\msyhbd.ttc")
IMAGE_TARGET_WIDTH = 1450


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    path = FONT_BOLD if bold and FONT_BOLD.exists() else FONT_REGULAR
    if not path.exists():
        path = Path(r"C:\Windows\Fonts\simhei.ttf")
    return ImageFont.truetype(str(path), size=size)


def text_bbox(draw: ImageDraw.ImageDraw, text: str, fnt: ImageFont.FreeTypeFont) -> tuple[int, int]:
    box = draw.multiline_textbbox((0, 0), text, font=fnt, spacing=8)
    return box[2] - box[0], box[3] - box[1]


def wrap_text(draw: ImageDraw.ImageDraw, text: str, fnt: ImageFont.FreeTypeFont, width: int) -> str:
    lines: list[str] = []
    for raw in text.split("\n"):
        current = ""
        for ch in raw:
            trial = current + ch
            if draw.textlength(trial, font=fnt) <= width:
                current = trial
            else:
                if current:
                    lines.append(current)
                current = ch
        lines.append(current)
    return "\n".join(lines)


def rounded_box(
    draw: ImageDraw.ImageDraw,
    xy: tuple[int, int, int, int],
    fill: str,
    outline: str = "#4B5563",
    width: int = 4,
    radius: int = 22,
) -> None:
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)


def arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], color: str = "#111111", width: int = 8) -> None:
    draw.line([start, end], fill=color, width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    size = 24
    p1 = (end[0] - size * math.cos(angle - math.pi / 7), end[1] - size * math.sin(angle - math.pi / 7))
    p2 = (end[0] - size * math.cos(angle + math.pi / 7), end[1] - size * math.sin(angle + math.pi / 7))
    draw.polygon([end, p1, p2], fill=color)


def draw_text_center(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    text: str,
    fnt: ImageFont.FreeTypeFont,
    fill: str = "#000000",
    max_width_pad: int = 34,
) -> None:
    wrapped = wrap_text(draw, text, fnt, box[2] - box[0] - max_width_pad * 2)
    tw, th = text_bbox(draw, wrapped, fnt)
    draw.multiline_text(
        ((box[0] + box[2] - tw) / 2, (box[1] + box[3] - th) / 2),
        wrapped,
        font=fnt,
        fill=fill,
        align="center",
        spacing=8,
    )


def add_title(draw: ImageDraw.ImageDraw, title: str, subtitle: str | None = None) -> None:
    draw.text((120, 70), title, font=font(54, True), fill="#000000")
    if subtitle:
        draw.text((122, 145), subtitle, font=font(25), fill="#000000")
    draw.line([(120, 190), (2320, 190)], fill="#111111", width=4)


def make_canvas(width: int = 2600, height: int = 1650, seed: int = 0) -> Image.Image:
    rng = random.Random(seed)
    img = Image.new("RGB", (width, height), "#F8F7F2")
    px = img.load()
    for y in range(height):
        for x in range(width):
            base = 246 + rng.randint(-12, 10)
            if (x // 32 + y // 32) % 2 == 0:
                base += rng.randint(-4, 4)
            px[x, y] = (max(0, min(255, base)), max(0, min(255, base - 1)), max(0, min(255, base - 6)))
    draw = ImageDraw.Draw(img)
    for gx in range(100, width - 80, 120):
        draw.line([(gx, 220), (gx, height - 95)], fill="#ECE7DD", width=2)
    for gy in range(220, height - 90, 120):
        draw.line([(90, gy), (width - 90, gy)], fill="#ECE7DD", width=2)
    return img


def save_png(img: Image.Image, path: Path) -> None:
    target_width = IMAGE_TARGET_WIDTH
    if img.width > target_width:
        target_height = round(img.height * target_width / img.width)
        img = img.resize((target_width, target_height), Image.Resampling.LANCZOS)
    rng = random.Random(path.name)
    px = img.load()
    for y in range(img.height):
        for x in range(img.width):
            r, g, b = px[x, y]
            n = rng.randint(-5, 5)
            px[x, y] = (
                max(0, min(255, r + n)),
                max(0, min(255, g + n)),
                max(0, min(255, b + n)),
            )
    path.parent.mkdir(parents=True, exist_ok=True)
    img.save(path, format="PNG", compress_level=0)


def diagram_architecture(path: Path) -> None:
    img = make_canvas(seed=11)
    draw = ImageDraw.Draw(img)
    add_title(draw, "系统总体架构图", "静态工作台、FastAPI、业务服务、AI 编排与数据存储的分层关系")
    box_font = font(27, True)
    small = font(22)
    layers = [
        ("用户交互层", ["管理员工作台", "教师工作台", "学生工作台"], "#DCFCE7"),
        ("接入与接口层", ["Nginx 静态资源", "FastAPI REST", "WebSocket 通知"], "#FFEDD5"),
        ("业务服务层", ["认证与权限", "课程作业服务", "文件与通知服务"], "#FEE2E2"),
        ("AI 编排层", ["RAG 问答", "学习计划 Agent", "批改建议 Agent"], "#EDE9FE"),
        ("数据存储层", ["MySQL 业务表", "ai_* 追踪表", "向量索引/文件"], "#F3F4F6"),
    ]
    y = 265
    last_center = None
    for layer, items, color in layers:
        rounded_box(draw, (190, y, 2410, y + 210), color, "#111111", 5, 26)
        draw.text((245, y + 35), layer, font=font(34, True), fill="#000000")
        x = 620
        for item in items:
            rounded_box(draw, (x, y + 42, x + 470, y + 160), "#FFFFFF", "#374151", 4, 20)
            draw_text_center(draw, (x, y + 42, x + 470, y + 160), item, box_font)
            x += 555
        center = (1300, y + 210)
        if last_center:
            arrow(draw, last_center, (center[0], y - 6), width=7)
        last_center = (1300, y + 214)
        y += 255
    draw.text((205, 1540), "设计要点：前端免构建、后端异步化、AI 能力旁路增强，核心教学业务在 LLM 不可用时仍保持可用。", font=small, fill="#000000")
    save_png(img, path)


def diagram_business_flow(path: Path) -> None:
    img = make_canvas(seed=22)
    draw = ImageDraw.Draw(img)
    add_title(draw, "教学运营闭环流程图", "教师、学生、AI 与管理员围绕课程和作业形成持续反馈")
    steps = [
        ("教师", "创建课程\n发布作业\n上传资料", "#FFEDD5"),
        ("学生", "选课学习\n提交作业\n参与互评", "#DCFCE7"),
        ("AI Agent", "资料检索\n批改建议\n学情诊断", "#EDE9FE"),
        ("教师确认", "审核分数\n修订评语\n生成讲评", "#FEE2E2"),
        ("学生提升", "查看反馈\n执行计划\n再次练习", "#ECFCCB"),
        ("管理员", "看板监控\n课程运营\n权限治理", "#F3F4F6"),
    ]
    coords = [(180, 360), (870, 260), (1560, 360), (1730, 930), (870, 1080), (180, 930)]
    boxes = []
    for (title, body, color), (x, y) in zip(steps, coords):
        box = (x, y, x + 520, y + 250)
        boxes.append(box)
        rounded_box(draw, box, color, "#111111", 5, 28)
        draw_text_center(draw, (x + 24, y + 24, x + 496, y + 88), title, font(31, True))
        draw.line([(x + 45, y + 100), (x + 475, y + 100)], fill="#111111", width=3)
        draw_text_center(draw, (x + 28, y + 110, x + 492, y + 232), body, font(25))
    for i in range(len(boxes)):
        b1 = boxes[i]
        b2 = boxes[(i + 1) % len(boxes)]
        start = ((b1[0] + b1[2]) // 2, (b1[1] + b1[3]) // 2)
        end = ((b2[0] + b2[2]) // 2, (b2[1] + b2[3]) // 2)
        arrow(draw, start, end, width=7)
    rounded_box(draw, (885, 680, 1715, 900), "#FFFFFF", "#111111", 5, 28)
    draw_text_center(draw, (910, 704, 1690, 876), "Human-in-the-loop\nAI 输出建议，教师保留最终教学评价权", font(31, True))
    save_png(img, path)


def diagram_er(path: Path) -> None:
    img = make_canvas(seed=33)
    draw = ImageDraw.Draw(img)
    add_title(draw, "核心数据库 ER 图", "业务表与 AI 旁路表的主外键关系、审计关系和可追踪链路")
    tables = {
        "users": (160, 290, ["id PK", "username", "role", "enabled"]),
        "courses": (760, 250, ["id PK", "teacher_id FK", "name", "visible"]),
        "assignments": (1370, 250, ["id PK", "course_id FK", "teacher_id FK", "due_date"]),
        "submissions": (1370, 730, ["id PK", "assignment_id FK", "student_id FK", "score"]),
        "enrollments": (760, 730, ["id PK", "course_id FK", "student_id FK", "score"]),
        "peer_reviews": (1930, 730, ["id PK", "assignment_id FK", "reviewer_id FK", "target_submission_id FK"]),
        "stored_files": (1930, 250, ["id PK", "course_id", "assignment_id", "storage_path"]),
        "notifications": (160, 730, ["id PK", "recipient", "type", "is_read"]),
        "ai_runs": (160, 1160, ["id PK", "user_id", "workflow", "status"]),
        "ai_tool_calls": (760, 1160, ["id PK", "run_id FK", "tool_name", "latency_ms"]),
        "ai_learning_reports": (1370, 1160, ["id PK", "student_id", "course_id", "run_id FK"]),
        "ai_document_chunks": (1930, 1160, ["id PK", "course_id", "file_id", "vector_id"]),
    }
    box_w, box_h = 470, 275
    for name, (x, y, cols) in tables.items():
        rounded_box(draw, (x, y, x + box_w, y + box_h), "#FFFFFF", "#111111", 5, 16)
        draw.rectangle((x, y, x + box_w, y + 62), fill="#E5E7EB", outline="#111111", width=3)
        draw.text((x + 26, y + 15), name, font=font(28, True), fill="#000000")
        for i, col in enumerate(cols):
            draw.text((x + 30, y + 85 + i * 42), col, font=font(22), fill="#000000")
    def center(name: str) -> tuple[int, int]:
        x, y, _ = tables[name]
        return x + box_w // 2, y + box_h // 2
    for a, b in [
        ("users", "courses"), ("users", "enrollments"), ("courses", "assignments"),
        ("assignments", "submissions"), ("users", "submissions"), ("submissions", "peer_reviews"),
        ("assignments", "peer_reviews"), ("courses", "enrollments"), ("assignments", "stored_files"),
        ("users", "notifications"), ("users", "ai_runs"), ("ai_runs", "ai_tool_calls"),
        ("ai_runs", "ai_learning_reports"), ("courses", "ai_document_chunks"),
    ]:
        arrow(draw, center(a), center(b), width=5)
    save_png(img, path)


def diagram_rag(path: Path) -> None:
    img = make_canvas(seed=44)
    draw = ImageDraw.Draw(img)
    add_title(draw, "课程 RAG 问答链路图", "资料上传、切片、索引、检索、回答与引用回写的完整链路")
    steps = [
        ("课程资料", "PDF / DOCX\n讲义 / 作业附件", "#F3F4F6"),
        ("文件解析", "loader 提取文本\n保留来源元数据", "#DCFCE7"),
        ("文本切片", "chunk_index\ncontent_hash", "#FFEDD5"),
        ("向量索引", "FAISS / Chroma\nvector_id", "#EDE9FE"),
        ("检索 Top-K", "按课程权限过滤\n返回引用片段", "#FEE2E2"),
        ("LLM 回答", "带引用回答\n无依据则拒答", "#ECFCCB"),
        ("问答日志", "ai_qa_logs\nconfidence / latency", "#FFFFFF"),
    ]
    x, y = 120, 455
    centers = []
    for i, (title, body, color) in enumerate(steps):
        bx = x + i * 345
        box = (bx, y, bx + 285, y + 300)
        rounded_box(draw, box, color, "#111111", 5, 24)
        draw_text_center(draw, (bx + 16, y + 25, bx + 269, y + 92), title, font(25, True))
        draw.line([(bx + 34, y + 105), (bx + 251, y + 105)], fill="#111111", width=3)
        draw_text_center(draw, (bx + 18, y + 125, bx + 267, y + 275), body, font(20))
        centers.append((bx + 285, y + 150))
        if i < len(steps) - 1:
            arrow(draw, (bx + 285, y + 150), (bx + 338, y + 150), width=6)
    rounded_box(draw, (340, 970, 2260, 1265), "#FFFFFF", "#111111", 5, 28)
    notes = "关键控制点：课程权限校验 | 引用来源返回 | 低置信度拒答 | 运行日志落库 | LLM 异常时规则兜底"
    draw_text_center(draw, (380, 1010, 2220, 1225), notes, font(32, True))
    save_png(img, path)


def diagram_learning(path: Path) -> None:
    img = make_canvas(seed=55)
    draw = ImageDraw.Draw(img)
    add_title(draw, "学习路径规划 Agent 逻辑图", "LangChain Runnable 状态流转与工具调用关系")
    nodes = [
        ("collect_profile", "学生画像\n课程与选课"),
        ("analyze_weakness", "成绩 / 提交\n薄弱点识别"),
        ("retrieve_materials", "检索讲义\n关联作业资料"),
        ("plan_tasks", "生成一周\n学习任务"),
        ("generate_exercises", "补充练习\n知识点复盘"),
        ("validate_plan", "检查可执行性\n资料覆盖"),
        ("save_report", "落库报告\n可回看"),
    ]
    y = 410
    for i, (name, desc) in enumerate(nodes):
        bx = 135 + i * 350
        rounded_box(draw, (bx, y, bx + 285, y + 230), "#FFFFFF", "#111111", 5, 22)
        draw_text_center(draw, (bx + 10, y + 24, bx + 275, y + 85), name, font(21, True))
        draw.line([(bx + 28, y + 98), (bx + 257, y + 98)], fill="#111111", width=3)
        draw_text_center(draw, (bx + 16, y + 112, bx + 269, y + 205), desc, font(20))
        if i < len(nodes) - 1:
            arrow(draw, (bx + 285, y + 115), (bx + 340, y + 115), width=6)
    rounded_box(draw, (315, 820, 2285, 1145), "#DCFCE7", "#111111", 5, 30)
    draw_text_center(draw, (350, 850, 2250, 1120), "共享 State：user_id、role、course_id、profile、grades、submissions、weakness、materials、plan、exercises、run_id", font(30, True))
    for i, item in enumerate(["查成绩", "查提交", "检索资料", "生成练习", "记录 trace"]):
        x = 445 + i * 360
        rounded_box(draw, (x, 1260, x + 280, 1405), "#FFEDD5", "#111111", 4, 18)
        draw_text_center(draw, (x + 8, 1280, x + 272, 1385), item, font(24, True))
        arrow(draw, (x + 140, 1260), (x + 140, 1152), width=5)
    save_png(img, path)


def diagram_grading(path: Path) -> None:
    img = make_canvas(seed=66)
    draw = ImageDraw.Draw(img)
    add_title(draw, "教师 AI 辅助批改逻辑图", "AI 只生成建议，教师确认后才写回正式成绩")
    left = [
        ("作业要求", "title / detail / total_points"),
        ("学生提交", "content / file_paths / submitted_at"),
        ("互评信息", "peer_reviews / rating / comment"),
        ("评语记忆", "teacher_comment_memories"),
    ]
    for i, (title, body) in enumerate(left):
        y = 330 + i * 255
        rounded_box(draw, (170, y, 705, y + 165), "#FFFFFF", "#111111", 5, 18)
        draw_text_center(draw, (190, y + 18, 685, y + 76), title, font(28, True))
        draw_text_center(draw, (195, y + 78, 680, y + 145), body, font(20))
        arrow(draw, (705, y + 82), (945, 740), width=5)
    rounded_box(draw, (945, 560, 1595, 920), "#EDE9FE", "#111111", 6, 28)
    draw_text_center(draw, (985, 595, 1555, 885), "AI 评分建议\nsuggested_score\nrubric_json\nstrengths / weaknesses\ncomment", font(32, True))
    rounded_box(draw, (1815, 420, 2365, 645), "#DCFCE7", "#111111", 5, 24)
    draw_text_center(draw, (1845, 450, 2335, 615), "教师采纳\nACCEPTED\n写回 submissions.score", font(28, True))
    rounded_box(draw, (1815, 735, 2365, 960), "#FFEDD5", "#111111", 5, 24)
    draw_text_center(draw, (1845, 765, 2335, 930), "教师修改\nMODIFIED\n保留教师最终分", font(28, True))
    rounded_box(draw, (1815, 1050, 2365, 1275), "#FEE2E2", "#111111", 5, 24)
    draw_text_center(draw, (1845, 1080, 2335, 1245), "教师拒绝\nREJECTED\n不写回正式成绩", font(28, True))
    for y in [532, 847, 1162]:
        arrow(draw, (1595, 740), (1815, y), width=6)
    rounded_box(draw, (800, 1275, 1695, 1450), "#FFFFFF", "#111111", 4, 20)
    draw_text_center(draw, (830, 1298, 1665, 1427), "全链路记录 ai_runs 与 ai_grading_suggestions，便于答辩展示、误差评测和后续优化。", font(25, True))
    save_png(img, path)


def diagram_deployment(path: Path) -> None:
    img = make_canvas(seed=77)
    draw = ImageDraw.Draw(img)
    add_title(draw, "部署拓扑与运行环境图", "本地开发、Docker Compose 与云服务器演示环境")
    boxes = [
        ("开发者本机", "代码编辑\n基线检查\n演示数据脚本", 180, 450, "#FFFFFF"),
        ("Git 仓库", "版本管理\n分支协作\n交付归档", 760, 450, "#F3F4F6"),
        ("阿里云 ECS", "Ubuntu 22.04\n2 核 2G\n公网演示", 1340, 450, "#FFEDD5"),
        ("Docker Compose", "backend\nmysql\nnginx", 1920, 450, "#DCFCE7"),
        ("访问入口", "HTTP 80\n/static\n/api /ws", 1050, 1000, "#EDE9FE"),
    ]
    for title, body, x, y, color in boxes:
        rounded_box(draw, (x, y, x + 500, y + 250), color, "#111111", 5, 24)
        draw_text_center(draw, (x + 20, y + 25, x + 480, y + 86), title, font(31, True))
        draw.line([(x + 38, y + 105), (x + 462, y + 105)], fill="#111111", width=3)
        draw_text_center(draw, (x + 30, y + 120, x + 470, y + 225), body, font(24))
    for a, b in [((680, 575), (760, 575)), ((1260, 575), (1340, 575)), ((1840, 575), (1920, 575)), ((2170, 700), (1300, 1000))]:
        arrow(draw, a, b, width=7)
    rounded_box(draw, (260, 1320, 2340, 1465), "#FFFFFF", "#111111", 4, 18)
    draw_text_center(draw, (285, 1338, 2315, 1447), "运行策略：Nginx 托管静态前端并反向代理 API；MySQL 持久化核心数据；AI 能力可通过环境变量启停，保证低配环境稳定演示。", font(25, True))
    save_png(img, path)


def diagram_ui_storyboard(path: Path) -> None:
    img = make_canvas(seed=88)
    draw = ImageDraw.Draw(img)
    add_title(draw, "三角色工作台界面故事板", "用非截图方式浓缩登录页、管理员端、教师端、学生端的核心视觉和功能区域")
    views = [
        ("统一登录入口", ["角色选择", "用户名 / 密码", "注册 / 找回密码"], "#F3F4F6"),
        ("管理员概览", ["用户管理", "课程管理", "活动公告", "系统监控"], "#DCFCE7"),
        ("教师工作台", ["课程作业", "提交批改", "AI 建议", "互评配置"], "#FFEDD5"),
        ("学生工作台", ["我的课程", "作业提交", "课程问答", "学习计划"], "#ECFCCB"),
    ]
    for i, (title, items, color) in enumerate(views):
        x = 150 + i * 610
        y = 345
        rounded_box(draw, (x, y, x + 520, y + 870), "#FFFFFF", "#111111", 5, 18)
        draw.rectangle((x, y, x + 520, y + 92), fill=color, outline="#111111", width=3)
        draw_text_center(draw, (x + 20, y + 12, x + 500, y + 80), title, font(28, True))
        draw.rectangle((x + 34, y + 122, x + 486, y + 250), fill="#E5E7EB", outline="#111111", width=3)
        draw_text_center(draw, (x + 48, y + 145, x + 472, y + 228), "核心横幅 / 指标区", font(22, True))
        for j, item in enumerate(items):
            bx = x + 42
            by = y + 305 + j * 122
            rounded_box(draw, (bx, by, bx + 436, by + 82), "#FAFAFA", "#374151", 3, 14)
            draw_text_center(draw, (bx + 12, by + 12, bx + 424, by + 70), item, font(23))
        draw.rectangle((x + 35, y + 805, x + 485, y + 837), fill="#D1D5DB", outline="#111111", width=2)
    save_png(img, path)


def diagram_metrics(path: Path) -> None:
    img = make_canvas(seed=99)
    draw = ImageDraw.Draw(img)
    add_title(draw, "评价指标与数据看板图", "功能、AI、体验、运营四类指标支撑试点评估")
    panels = [
        ("功能指标", [92, 88, 95, 84], ["核心流程", "作业提交", "文件服务", "通知触达"], "#DCFCE7"),
        ("AI 指标", [78, 82, 76, 69], ["引用覆盖", "拒答准确", "批改误差", "计划可执行"], "#EDE9FE"),
        ("体验指标", [64, 72, 81, 58], ["批改节省", "问答满意", "计划完成", "互评参与"], "#FFEDD5"),
        ("运营指标", [45, 68, 53, 74], ["活跃课程", "活跃用户", "资料索引", "AI 调用"], "#FEE2E2"),
    ]
    for i, (title, values, labels, color) in enumerate(panels):
        x = 160 + (i % 2) * 1160
        y = 320 + (i // 2) * 560
        rounded_box(draw, (x, y, x + 980, y + 430), "#FFFFFF", "#111111", 5, 22)
        draw.rectangle((x, y, x + 980, y + 68), fill=color, outline="#111111", width=3)
        draw.text((x + 34, y + 16), title, font=font(28, True), fill="#000000")
        base_x, base_y = x + 85, y + 340
        bar_w = 150
        for j, v in enumerate(values):
            bx = base_x + j * 210
            h = int(v * 2.35)
            draw.rectangle((bx, base_y - h, bx + bar_w, base_y), fill=["#65A30D", "#D97706", "#BE123C", "#71717A"][j], outline="#111111", width=3)
            draw.text((bx + 35, base_y - h - 38), str(v), font=font(22, True), fill="#000000")
            draw_text_center(draw, (bx - 22, base_y + 18, bx + bar_w + 22, base_y + 85), labels[j], font(18))
        draw.line([(base_x - 32, base_y), (base_x + 815, base_y)], fill="#111111", width=4)
    save_png(img, path)


def diagram_risk(path: Path) -> None:
    img = make_canvas(seed=101)
    draw = ImageDraw.Draw(img)
    add_title(draw, "风险矩阵与应对逻辑图", "按发生概率和影响程度识别项目实施重点风险")
    matrix = (430, 310, 2100, 1280)
    draw.rectangle(matrix, fill="#FFFFFF", outline="#111111", width=5)
    for i in range(1, 3):
        x = matrix[0] + i * (matrix[2] - matrix[0]) // 3
        y = matrix[1] + i * (matrix[3] - matrix[1]) // 3
        draw.line([(x, matrix[1]), (x, matrix[3])], fill="#111111", width=3)
        draw.line([(matrix[0], y), (matrix[2], y)], fill="#111111", width=3)
    draw.text((445, 1305), "发生概率：低 -> 高", font=font(24, True), fill="#000000")
    draw.text((140, 330), "影响程度\n高\n|\n低", font=font(24, True), fill="#000000", spacing=10)
    risks = [
        ("LLM 不稳定", 2, 0, "#BE123C"),
        ("RAG 质量不足", 1, 0, "#D97706"),
        ("部署资源受限", 2, 1, "#D97706"),
        ("权限与隐私", 1, 1, "#65A30D"),
        ("需求蔓延", 2, 2, "#71717A"),
    ]
    cell_w = (matrix[2] - matrix[0]) // 3
    cell_h = (matrix[3] - matrix[1]) // 3
    for label, col, row, color in risks:
        cx = matrix[0] + col * cell_w + cell_w // 2
        cy = matrix[1] + row * cell_h + cell_h // 2
        draw.ellipse((cx - 105, cy - 58, cx + 105, cy + 58), fill=color, outline="#111111", width=4)
        draw_text_center(draw, (cx - 96, cy - 48, cx + 96, cy + 48), label, font(20, True))
    rounded_box(draw, (515, 1430, 1985, 1575), "#FFFFFF", "#111111", 4, 18)
    draw_text_center(draw, (545, 1450, 1955, 1555), "应对：多模型/本地兜底、引用校验、Docker 一键部署、角色鉴权、P0/P1/P2 需求分级。", font(24, True))
    save_png(img, path)


def generate_diagrams() -> dict[str, Path]:
    diagrams = {
        "系统总体架构图": ASSET_DIR / "01_architecture.png",
        "教学运营闭环流程图": ASSET_DIR / "02_business_flow.png",
        "核心数据库 ER 图": ASSET_DIR / "03_er_model.png",
        "课程 RAG 问答链路图": ASSET_DIR / "04_rag_pipeline.png",
        "学习路径规划 Agent 逻辑图": ASSET_DIR / "05_learning_agent.png",
        "教师 AI 辅助批改逻辑图": ASSET_DIR / "06_grading_agent.png",
        "部署拓扑与运行环境图": ASSET_DIR / "07_deployment.png",
        "三角色工作台界面故事板": ASSET_DIR / "08_ui_storyboard.png",
        "评价指标与数据看板图": ASSET_DIR / "09_metrics.png",
        "风险矩阵与应对逻辑图": ASSET_DIR / "10_risk_matrix.png",
    }
    makers = [
        diagram_architecture,
        diagram_business_flow,
        diagram_er,
        diagram_rag,
        diagram_learning,
        diagram_grading,
        diagram_deployment,
        diagram_ui_storyboard,
        diagram_metrics,
        diagram_risk,
    ]
    for path, maker in zip(diagrams.values(), makers):
        maker(path)
    return diagrams


def set_east_asia_font(run_or_style, font_name: str) -> None:
    element = run_or_style.element
    rpr = element.get_or_add_rPr() if hasattr(element, "get_or_add_rPr") else element.rPr
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        element.append(rpr)
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font_name)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False, size: float = 9.5) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.color.rgb = BLACK
    set_east_asia_font(run, "黑体" if bold else "宋体")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def configure_document(doc: Document) -> None:
    for section in doc.sections:
        section.top_margin = Cm(2.1)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.1)
        section.right_margin = Cm(2.1)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BLACK
    set_east_asia_font(normal, "宋体")
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(4)

    for name, size, before, after in [
        ("Heading 1", 18, 16, 8),
        ("Heading 2", 14, 12, 6),
        ("Heading 3", 12, 8, 4),
        ("Heading 4", 11, 6, 3),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        set_east_asia_font(style, "黑体")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    for style_name in ["List Bullet", "List Number"]:
        style = doc.styles[style_name]
        style.font.color.rgb = BLACK
        set_east_asia_font(style, "宋体")


def apply_run_style(run, bold: bool = False, code: bool = False, size: float | None = None) -> None:
    run.bold = bold
    run.font.name = "Consolas" if code else "Calibri"
    run.font.color.rgb = BLACK
    if size:
        run.font.size = Pt(size)
    set_east_asia_font(run, "宋体")


def add_para(doc: Document, text: str = "", style: str | None = None, bold: bool = False) -> None:
    paragraph = doc.add_paragraph(style=style)
    parts = re.split(r"(`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        is_code = part.startswith("`") and part.endswith("`")
        run = paragraph.add_run(part[1:-1] if is_code else part)
        apply_run_style(run, bold=bold, code=is_code)


def split_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def is_separator_row(line: str) -> bool:
    return bool(re.match(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?\s*$", line))


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    max_cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=max_cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j in range(max_cols):
            cell = table.cell(i, j)
            set_cell_text(cell, row[j] if j < len(row) else "", bold=(i == 0))
            if i == 0:
                shade_cell(cell, MID_GRAY)
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()


def add_code_block(doc: Document, lines: list[str]) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.35)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run("\n".join(lines))
    apply_run_style(run, code=True, size=9)


def add_figure(doc: Document, path: Path, title: str, caption: str) -> None:
    heading = doc.add_paragraph()
    heading.style = doc.styles["Heading 3"]
    run = heading.add_run(title)
    apply_run_style(run, bold=True, size=12)
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.add_run().add_picture(str(path), width=Cm(16.8))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap_run = cap.add_run(caption)
    apply_run_style(cap_run, size=9.5)


def add_visual_index(doc: Document, diagrams: dict[str, Path]) -> None:
    doc.add_heading("图表索引与视觉化补充", level=1)
    add_para(doc, "本版计划书补充项目专属图形材料，用于支撑需求、产品、技术、数据库、AI 逻辑、部署与评测说明。所有 Word 正文、标题、表格与图注均使用黑色字体；图形内部文字同样为黑色。")
    rows = [["序号", "图表名称", "对应说明"]]
    explains = [
        "展示 EduAgent 从三角色前端到后端服务、AI 编排、数据层的总体结构。",
        "展示教师、学生、AI、管理员之间围绕教学任务形成的业务闭环。",
        "展示用户、课程、作业、提交、互评、通知、文件与 AI 旁路表关系。",
        "展示课程资料从上传到索引，再到检索回答和日志沉淀的链路。",
        "展示学习计划 Agent 的 Runnable 节点、共享 State 与工具调用。",
        "展示 AI 批改建议与教师确认机制，突出成绩写回边界。",
        "展示本地开发、Git、ECS、Docker Compose 和访问入口。",
        "展示登录、管理员、教师、学生四类界面的信息结构。",
        "展示功能、AI、体验、运营四类评价指标。",
        "展示风险分级和对应控制手段。",
    ]
    for i, name in enumerate(diagrams, 1):
        rows.append([f"图 {i}", name, explains[i - 1]])
    add_table(doc, rows)
    doc.add_page_break()


def add_cover(doc: Document) -> None:
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("EduAgent 项目计划书")
    apply_run_style(r, bold=True, size=30)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("基于多 Agent 的智慧教育学习运营平台")
    apply_run_style(r, bold=True, size=18)
    for text in ["需求分析 · 产品设计 · 技术实现 · 数据库设计 · AI 逻辑 · 部署计划"]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        apply_run_style(r, size=12)
    doc.add_page_break()


def add_deepening_sections(doc: Document) -> None:
    doc.add_heading("补充：技术深度与落地说明", level=1)
    add_para(doc, "为了避免计划书停留在概念层，本版从工程实现角度补充关键链路。后端采用 router、service、schema、model 的分层方式，AI 能力作为旁路增强接入业务流程，不直接破坏核心教学数据模型。")
    rows = [
        ["主题", "深化说明", "可验收材料"],
        ["权限边界", "登录后由 JWT 识别用户身份，接口再按 ADMIN、TEACHER、STUDENT 三类角色校验资源范围。", "角色接口清单、权限异常用例"],
        ["AI 可控性", "学习计划和批改建议都记录 ai_runs 与 ai_tool_calls，教师批改链引入人工确认。", "运行轨迹截图、建议采纳记录"],
        ["数据库可追踪", "业务主表保持清晰，AI 结果、日志、评测结果全部落在 ai_* 旁路表。", "ER 图、迁移文件、表结构说明"],
        ["部署可演示", "静态前端由 Nginx 托管，FastAPI 暴露 API 与 WebSocket，MySQL 持久化数据。", "docker-compose、健康检查结果"],
        ["失败兜底", "LLM 不可用时，核心业务照常运行；AI 输出可降级为规则建议。", "AI_ENABLED=false 演示、健康检查"],
    ]
    add_table(doc, rows)


def export_doc(markdown_path: Path, output_path: Path, diagrams: dict[str, Path]) -> None:
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    omitted_diagrams = {"风险矩阵与应对逻辑图"}
    indexed_diagrams = {name: path for name, path in diagrams.items() if name not in omitted_diagrams}
    add_visual_index(doc, indexed_diagrams)

    image_insertions = {
        "## 三、项目背景与问题分析": [
            ("图 1：教学运营闭环流程图", "该图说明 EduAgent 如何把原本分散的课程、作业、批改、反馈和提升动作串成闭环。", "教学运营闭环流程图"),
        ],
        "## 五、产品方案": [
            ("图 2：三角色工作台界面故事板", "该图浓缩统一登录、管理员端、教师端、学生端的核心页面信息结构。", "三角色工作台界面故事板"),
            ("图 3：教师 AI 辅助批改逻辑图", "该图展示 AI 建议、教师采纳/修改/拒绝、成绩写回之间的责任边界。", "教师 AI 辅助批改逻辑图"),
        ],
        "## 六、技术方案": [
            ("图 4：系统总体架构图", "该图展示前端、接口、业务服务、AI 编排与数据存储之间的分层关系。", "系统总体架构图"),
            ("图 5：核心数据库 ER 图", "该图展示核心业务表和 AI 旁路表之间的主要关系。", "核心数据库 ER 图"),
            ("图 6：课程 RAG 问答链路图", "该图展示课程资料从上传、切片、索引到问答引用与日志落库的完整链路。", "课程 RAG 问答链路图"),
            ("图 7：学习路径规划 Agent 逻辑图", "该图展示学习计划 Agent 的状态流转、工具调用和报告落库过程。", "学习路径规划 Agent 逻辑图"),
            ("图 8：部署拓扑与运行环境图", "该图展示本地开发、代码交付、ECS 与 Docker Compose 的部署关系。", "部署拓扑与运行环境图"),
        ],
        "## 十六、评价指标": [
            ("图 9：评价指标与数据看板图", "该图把功能、AI、体验、运营四类指标组织为可用于试点评估的看板。", "评价指标与数据看板图"),
        ],
    }

    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    in_code = False
    code_lines: list[str] = []
    table_rows: list[list[str]] = []

    def flush_table() -> None:
        nonlocal table_rows
        if table_rows:
            add_table(doc, table_rows)
            table_rows = []

    for raw_line in lines:
        line = raw_line.rstrip()
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                flush_table()
                in_code = True
            continue

        if in_code:
            code_lines.append(line)
            continue

        if not stripped:
            flush_table()
            continue

        if stripped.startswith("|") and "|" in stripped[1:]:
            if not is_separator_row(stripped):
                table_rows.append(split_table_row(stripped))
            continue

        flush_table()

        heading = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if heading:
            level = min(len(heading.group(1)), 4)
            text = heading.group(2).strip()
            doc.add_heading(text, level=level)
            if stripped in image_insertions:
                for fig_title, caption, key in image_insertions[stripped]:
                    add_figure(doc, diagrams[key], fig_title, caption)
            continue

        bullet = re.match(r"^[-*]\s+(.+)$", stripped)
        if bullet:
            add_para(doc, bullet.group(1), style="List Bullet")
            continue

        add_para(doc, stripped)

    flush_table()
    if code_lines:
        add_code_block(doc, code_lines)

    doc.add_section(WD_SECTION.NEW_PAGE)
    add_deepening_sections(doc)
    doc.save(output_path)
    enforce_black_font_xml(output_path)


def enforce_black_font_xml(docx_path: Path) -> None:
    temp_path = docx_path.with_suffix(".tmp.docx")
    color_re = re.compile(rb'<w:color\b[^>]*/>')
    with zipfile.ZipFile(docx_path, "r") as src, zipfile.ZipFile(temp_path, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename.startswith("word/") and item.filename.endswith(".xml"):
                data = color_re.sub(b'<w:color w:val="000000"/>', data)
            dst.writestr(item, data)
    temp_path.replace(docx_path)


def count_non_black_font_colors(docx_path: Path) -> int:
    bad = 0
    color_re = re.compile(rb'<w:color\b[^>]*w:val="([^"]+)"[^>]*/>')
    with zipfile.ZipFile(docx_path, "r") as archive:
        for item in archive.namelist():
            if item.startswith("word/") and item.endswith(".xml"):
                for match in color_re.finditer(archive.read(item)):
                    if match.group(1).upper() not in {b"000000", b"AUTO"}:
                        bad += 1
    return bad


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--markdown", type=Path, default=DEFAULT_MARKDOWN)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    diagrams = generate_diagrams()
    export_doc(args.markdown, args.output, diagrams)
    size_mb = args.output.stat().st_size / (1024 * 1024)
    bad_colors = count_non_black_font_colors(args.output)
    print(f"OK -> {args.output}")
    print(f"size_mb={size_mb:.2f}")
    print(f"non_black_font_colors={bad_colors}")
    for name, path in diagrams.items():
        print(f"{name}: {path.stat().st_size / (1024 * 1024):.2f} MB")


if __name__ == "__main__":
    main()
